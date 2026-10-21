"""Bounded CSV/Excel readers. Workbooks are never executed or saved server-side."""
import csv
import io
import json
import unicodedata
import zipfile
from datetime import date, datetime

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10000
MAX_COLUMNS = 100
MAX_EXPANDED_BYTES = 50 * 1024 * 1024


class CollectionValidationError(ValueError):
    pass


ALIASES = {
    "content": ("content", "review", "text", "reviewcontent", "nội dung", "bình luận", "đánh giá", "nhận xét"),
    "review_id": ("review_id", "reviewid", "reviewrateid", "reviewrelid", "id", "mã đánh giá"),
    "rating": ("rating", "stars", "star", "số sao", "sao"),
    "customer_name": ("customer_name", "buyername", "buyer_name", "tên khách hàng", "người mua"),
    "created_at": ("created_at", "review_date", "date", "reviewtime", "ngày đánh giá", "ngày"),
}
HEADERS = {alias: canonical for canonical, aliases in ALIASES.items() for alias in aliases}


def headers_for(values):
    headers = []
    for value in values:
        name = unicodedata.normalize("NFC", str(value or "")).strip().lstrip("\ufeff")
        headers.append(HEADERS.get(name.casefold(), name))
    if len(headers) > MAX_COLUMNS:
        raise CollectionValidationError("File tối đa 100 cột.")
    nonempty = [name for name in headers if name]
    if len(nonempty) != len(set(nonempty)) or "content" not in nonempty:
        raise CollectionValidationError("File cần cột content (hoặc Nội dung / Bình luận / Review), không trùng tên cột.")
    return headers


def row_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def parse_excel(raw, file_format):
    def rows_from_values(values):
        iterator = iter(values)
        try:
            header = headers_for(next(iterator))
        except StopIteration:
            raise CollectionValidationError("File Excel không có dữ liệu.")
        rows = []
        for index, values in enumerate(iterator, 2):
            if index > MAX_ROWS + 1:
                raise CollectionValidationError("Tối đa 10.000 dòng mỗi file; hãy chia file để gộp nhiều lần.")
            if not any(value is not None and value != "" for value in values):
                continue
            if len(values) > len(header) and any(value is not None and value != "" for value in values[len(header):]):
                raise CollectionValidationError(f"Excel dòng {index}: có dữ liệu ngoài cột tiêu đề.")
            row = {key: row_value(value) for key, value in zip(header, values) if key}
            # Excel numbers beyond this range can already have lost precision.
            if isinstance(row.get("review_id"), int) and abs(row["review_id"]) >= 2**53:
                raise CollectionValidationError(f"Excel dòng {index}: lưu review_id dài dưới dạng Text để giữ chính xác.")
            rows.append(row)
        return rows

    if file_format == "xlsx":
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise CollectionValidationError("Nội dung Excel giải nén quá lớn. Hãy chia file.")
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = workbook.worksheets[0]
            # Ignore misleading worksheet dimensions; cap actual streamed rows.
            sheet.reset_dimensions()
            def values():
                for cells in sheet.iter_rows(max_col=MAX_COLUMNS + 1):
                    row = []
                    for cell in cells:
                        if cell.data_type == "f":
                            raise CollectionValidationError("Excel có công thức. Hãy dùng Paste Special → Values trước khi tải lên.")
                        row.append(cell.value)
                    while row and row[-1] is None: row.pop()
                    yield row
            return rows_from_values(values())
        finally:
            workbook.close()
    import xlrd
    workbook = xlrd.open_workbook(file_contents=raw, on_demand=True, logfile=io.StringIO())
    try:
        sheet = workbook.sheet_by_index(0)
        if sheet.ncols > MAX_COLUMNS or sheet.nrows > MAX_ROWS + 1:
            raise CollectionValidationError("Excel tối đa 100 cột và 10.000 dòng dữ liệu.")
        def values():
            for row in sheet.get_rows():
                yield [xlrd.xldate_as_datetime(cell.value, workbook.datemode) if cell.ctype == xlrd.XL_CELL_DATE
                       else bool(cell.value) if cell.ctype == xlrd.XL_CELL_BOOLEAN else cell.value for cell in row]
        return rows_from_values(values())
    finally:
        workbook.release_resources()


def parse_file(raw, file_format):
    if len(raw) > MAX_BYTES:
        raise CollectionValidationError("File vượt quá 5 MiB; hãy chia thành các file nhỏ hơn.")
    try:
        if file_format in {"xlsx", "xls"}:
            return parse_excel(raw, file_format), {}
        text = raw.decode("utf-8-sig")
        if file_format == "json":  # Kept for existing API clients; UI uses CSV/Excel.
            data = json.loads(text)
            if isinstance(data, list): return data, {}
            if isinstance(data, dict) and isinstance(data.get("reviews"), list): return data["reviews"], data
            raise CollectionValidationError("JSON phải là mảng review hoặc object có mảng reviews.")
        if file_format != "csv":
            raise CollectionValidationError("Chỉ hỗ trợ CSV, Excel (.xlsx, .xls) hoặc JSON qua API.")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect=dialect, strict=True)
        try:
            header = headers_for(next(reader))
        except StopIteration:
            raise CollectionValidationError("File không có dữ liệu.")
        rows = []
        for values in reader:
            if not values: continue
            if len(values) != len(header):
                raise CollectionValidationError(f"CSV dòng {reader.line_num}: số cột không khớp header.")
            rows.append({key: value for key, value in zip(header, values) if key})
            if len(rows) > MAX_ROWS:
                raise CollectionValidationError("Tối đa 10.000 dòng mỗi file; hãy chia file để gộp nhiều lần.")
        return rows, {}
    except CollectionValidationError:
        raise
    except UnicodeDecodeError:
        raise CollectionValidationError("CSV/JSON phải dùng mã hóa UTF-8.")
    except Exception as error:
        # Never expose workbook contents, filenames or internal exception details.
        raise CollectionValidationError("File không đúng định dạng hoặc đã bị hỏng. Hãy lưu lại thành CSV / Excel hợp lệ.") from error
